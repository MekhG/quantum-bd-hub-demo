"""Weeks 5–6: knowledge base (documents, scripts, objections) and procurement & quotes."""

from __future__ import annotations

import io
from datetime import date, datetime, timedelta
from decimal import Decimal

import pandas as pd
import pytest
from docx import Document as DocxDocument
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import (
    Activity,
    ActivityType,
    CadenceStep,
    Contact,
    Document,
    Lead,
    Opportunity,
    Organisation,
    PriceItem,
    ProcurementItem,
    Product,
    Quote,
    Stage,
    Track,
    User,
)
from app.seed.seed_config import seed_all
from app.services.activity_log import LogEntry, log_activity
from app.services.cadence import IST, start_cadence, today_ist
from app.services.library import (
    DocumentData,
    LibraryError,
    ObjectionData,
    ScriptData,
    add_document,
    add_script,
    approve_document,
    approve_script,
    content_flags,
    documents_for_org,
    documents_frame,
    new_document_version,
    objections_frame,
    render,
    save_objection,
    search_library,
    set_objection_status,
    split_name,
    update_document,
)
from app.services.money import amount_in_words, inr
from app.services.procurement import (
    DealData,
    LineIn,
    ProcurementError,
    QuoteDraft,
    accept_quote,
    compute_totals,
    create_deal,
    create_quote,
    expire_quotes,
    fy_label,
    institution_code,
    lines_from_price_items,
    mark_quote_sent,
    pipeline_value,
    quote_docx,
    quotes_frame,
    revise_quote,
    save_price_items,
    suggest_route,
    update_deal,
)
from app.services.tasks import email_batch_frame


def _user(session: Session, name: str) -> User:
    return session.scalar(select(User).where(User.name.like(f"{name}%")))


def _lead(session: Session, name: str = "Omega University", *, track: str = "cold_colleges",
          channel: str = "Direct", management: str | None = None) -> Lead:
    org = Organisation(name=name, org_type="university", state="Kerala", management=management)
    session.add(org)
    session.flush()
    contact = Contact(org_id=org.id, name="Dr. Asha Menon", designation="Vice Chancellor",
                      email=f"vc@{name.split()[0].lower()}-example.ac.in", verified=True)
    session.add(contact)
    session.flush()
    lead = Lead(org_id=org.id, track_id=session.scalar(select(Track.id).where(Track.code == track)),
                stage=session.scalar(select(Stage).where(Stage.name == "New")), channel=channel,
                primary_contact_id=contact.id, owner_user_id=_user(session, "Riya").id)
    session.add(lead)
    session.commit()
    return lead


def _doc(**kw) -> DocumentData:
    data = dict(title="QSim Cluster brochure", doc_type="Brochure", url="https://drive.google.com/file/d/abc123/view",
                key_text="24-qubit state-vector simulation; Demo Notebook Studio notebooks", tags="qsim, brochure")
    data.update(kw)
    return DocumentData(**data)


# ---------------------------------------------------------------- knowledge base


def test_document_versions_and_approval(seeded: Session) -> None:
    riya, kabir = _user(seeded, "Riya"), _user(seeded, "Kabir")
    doc, flags = add_document(seeded, _doc(), kabir)
    assert doc.status == "draft" and doc.is_current and not flags
    with pytest.raises(PermissionError):
        approve_document(seeded, doc.id, kabir)  # only admins approve
    approve_document(seeded, doc.id, riya)
    with pytest.raises(LibraryError):
        update_document(seeded, doc.id, _doc(title="Changed"), kabir)  # approved → new version instead
    v2, _ = new_document_version(seeded, doc.id, _doc(url="https://drive.google.com/file/d/v2/view"), kabir)
    seeded.commit()
    # The approved v1 stays current until v2 is approved.
    assert seeded.get(Document, doc.id).is_current and not v2.is_current and v2.version_no == 2
    approve_document(seeded, v2.id, riya)
    seeded.commit()
    assert not seeded.get(Document, doc.id).is_current and seeded.get(Document, v2.id).is_current
    frame = documents_frame(seeded)
    assert list(frame["Version"]) == ["v2"]
    assert len(documents_frame(seeded, include_old=True)) == 2


def test_document_link_must_be_valid_and_unique(seeded: Session) -> None:
    kabir = _user(seeded, "Kabir")
    with pytest.raises(LibraryError):
        add_document(seeded, _doc(url="javascript:alert(1)"), kabir)
    add_document(seeded, _doc(), kabir)
    with pytest.raises(LibraryError):
        add_document(seeded, _doc(title="Same link"), kabir)


def test_search_finds_text_inside_key_text(seeded: Session) -> None:
    add_document(seeded, _doc(), _user(seeded, "Kabir"))
    seeded.commit()
    assert len(documents_frame(seeded, text="state-vector qsim")) == 1
    assert documents_frame(seeded, text="state-vector kits").empty
    assert search_library(seeded, "brochure")[0][0] == "document"


def test_content_check_flags_and_blocks_approval(seeded: Session) -> None:
    riya = _user(seeded, "Riya")
    assert content_flags(seeded, "Notebooks run on JupyterHub") == [("JupyterHub", "Demo Notebook Studio")]
    assert content_flags(seeded, "a plain crystal and a beam splitter") == []
    doc, flags = add_document(seeded, _doc(key_text="Runs on JupyterHub"), riya)
    assert flags
    with pytest.raises(LibraryError, match="Demo Notebook Studio"):
        approve_document(seeded, doc.id, riya)
    approve_document(seeded, doc.id, riya, override_flags=True)  # admin can keep it, and it's audited
    assert doc.status == "approved"


def test_scripts_fill_fields_and_feed_the_email_batch(seeded: Session) -> None:
    riya = _user(seeded, "Riya")
    lead = _lead(seeded)
    step = seeded.scalar(select(CadenceStep).where(CadenceStep.track_id == lead.track_id, CadenceStep.step_no == 1))
    with pytest.raises(LibraryError, match="first_nam"):
        add_script(seeded, ScriptData(title="Bad", kind="email", body="Hi {first_nam}"), riya)
    script, _ = add_script(seeded, ScriptData(
        title="Fresh — initial", kind="email", track_id=lead.track_id, cadence_step_id=step.id,
        subject="Quantum lab for {institution}", body="Dear {salutation} {last_name},\nRegards, {sender_name}"),
        riya)
    approve_script(seeded, script.id, riya)
    start_cadence(seeded, [lead.id], riya)
    seeded.commit()
    batch = email_batch_frame(seeded, None)
    row = batch[batch["Lead ID"] == lead.id].iloc[0]
    assert row["Script"] == "Fresh — initial"
    assert row["Subject"] == "Quantum lab for Omega University"
    assert row["Body"].startswith("Dear Dr Menon,") and "Riya Sharma" in row["Body"]


def test_name_splitting_and_render() -> None:
    assert split_name("Dr. Ananya Rao") == ("Dr", "Ananya", "Rao")
    assert split_name("Prof Kiran") == ("Prof.", "Kiran", "")
    text, empty = render("Dear {salutation} {first_name},", {"salutation": "", "first_name": "Asha"})
    assert text == "Dear Asha," and empty == ["salutation"]


def test_logging_shared_documents(seeded: Session) -> None:
    riya = _user(seeded, "Riya")
    lead = _lead(seeded)
    doc, _ = add_document(seeded, _doc(), riya)
    approve_document(seeded, doc.id, riya)
    activity, messages = log_activity(seeded, LogEntry(org_id=lead.org_id, lead_id=lead.id, type_name="Email",
                                                       when=datetime.now(IST), document_ids=[doc.id]), riya)
    seeded.commit()
    shared = documents_for_org(seeded, lead.org_id)
    assert list(shared["Document"]) == ["QSim Cluster brochure (v1)"]
    assert lead.stage.name == "Sent Brochure"  # not in a cadence, so sharing a brochure moves it on
    assert any("Sent Brochure" in m for m in messages)


def test_objections(seeded: Session) -> None:
    riya, kabir = _user(seeded, "Riya"), _user(seeded, "Kabir")
    assert len(objections_frame(seeded)) >= 5  # starter drafts from the seed
    row, _ = save_objection(seeded, ObjectionData("Too expensive", "Starting with one licence spreads the cost.",
                                                  tags="budget"), kabir)
    with pytest.raises(PermissionError):
        set_objection_status(seeded, row.id, "approved", kabir)
    set_objection_status(seeded, row.id, "approved", riya)
    seeded.commit()
    assert "Too expensive" in set(objections_frame(seeded, text="one licence")["They say"])


# ---------------------------------------------------------------- money


def test_indian_money_format_and_words() -> None:
    assert inr(Decimal("1100000")) == "₹11,00,000.00"
    assert inr(Decimal("8250000"), decimals=False) == "₹82,50,000"
    assert inr(999) == "₹999.00"
    assert amount_in_words(Decimal("2300000")) == "Rupees Twenty-Three Lakh Only"
    assert amount_in_words(Decimal("12345678.50")) == (
        "Rupees One Crore Twenty-Three Lakh Forty-Five Thousand Six Hundred Seventy-Eight and Fifty Paise Only")


def test_fy_and_codes() -> None:
    assert fy_label(date(2026, 10, 20)) == "2026-27"
    assert fy_label(date(2027, 3, 31)) == "2026-27"
    assert fy_label(date(2027, 4, 1)) == "2027-28"
    assert institution_code("Aranya University") == "AU"
    assert suggest_route("State Government") == "PAC (proprietary article)"  # not on GeM
    assert suggest_route("Private Un-Aided") == "Direct purchase order"


# ---------------------------------------------------------------- price list, deals, checklists


def _prices(session: Session, admin: User) -> None:
    frame = pd.DataFrame([
        {"Name": "NetSim Q — Quantum Network Simulator for Academia", "Product": "NetSim Q", "Unit": "licence",
         "Institute price (INR)": 1100000, "Partner price (INR)": 900000, "GST %": 18},
        {"Name": "Quantum Sensing Kit — teaching bundle", "Product":
            "Quantum Sensing Kit", "Unit": "kit", "Institute price (INR)": "24,50,000",
         "Partner price (INR)": None, "GST %": 18},
    ])
    save_price_items(session, frame, admin)
    session.commit()


def test_price_list_import_and_validation(seeded: Session) -> None:
    riya, kabir = _user(seeded, "Riya"), _user(seeded, "Kabir")
    with pytest.raises(PermissionError):
        save_price_items(seeded, pd.DataFrame([{"Name": "x"}]), kabir)
    _prices(seeded, riya)
    items = list(seeded.scalars(select(PriceItem).order_by(PriceItem.id)))
    assert [float(i.institute_price_inr) for i in items] == [1100000.0, 2450000.0]
    # Same name again updates instead of duplicating.
    added, updated = save_price_items(seeded, pd.DataFrame([{"Name": items[0].name, "Product": "NetSim Q",
                                                              "Institute price (INR)": 1150000, "GST %": 18}]), riya)
    assert (added, updated) == (0, 1)
    with pytest.raises(ProcurementError, match="unknown product"):
        save_price_items(seeded, pd.DataFrame([{"Name": "Y", "Product": "Nope"}]), riya)


def test_deal_checklist_follows_route_and_reseller(seeded: Session) -> None:
    riya = _user(seeded, "Riya")
    lead = _lead(seeded, channel="Reseller")
    deal = create_deal(seeded, lead.id, DealData(title="Phase 1", buying_route="PAC (proprietary article)",
                                                 expected_close=date(2027, 2, 15)), riya)
    seeded.commit()
    items = [i.item for i in seeded.scalars(select(ProcurementItem).where(ProcurementItem.opportunity_id == deal.id))]
    assert "GeM non-availability certificate (from the institution)" in items
    assert "Purchase order received" in items and "Advance payment received (as per payment terms)" in items
    assert "Fall-clause certificate from the reseller" in items  # reseller deal
    assert deal.budget_window_flag  # closes in Jan–Mar
    update_deal(seeded, deal.id, DealData(title="Phase 1", buying_route="Direct purchase order"), riya)
    seeded.commit()
    items = [i.item for i in seeded.scalars(select(ProcurementItem).where(ProcurementItem.opportunity_id == deal.id))]
    assert "Purchase order received" in items
    assert "GeM non-availability certificate (from the institution)" not in items  # untouched PAC items dropped


# ---------------------------------------------------------------- quotes


def test_totals_with_discount_and_gst() -> None:
    lines = [LineIn("A", 1, Decimal("1100000")), LineIn("B", 2, Decimal("100000"), gst_pct=Decimal("12"))]
    t = compute_totals(lines, Decimal("10"))
    assert t.subtotal == Decimal("1300000.00") and t.discount == Decimal("130000.00")
    assert t.taxable == Decimal("1170000.00")
    assert t.gst == Decimal("178200.00") + Decimal("21600.00")
    assert t.total == t.taxable + t.gst


def test_quote_lifecycle(seeded: Session) -> None:
    riya = _user(seeded, "Riya")
    _prices(seeded, riya)
    lead = _lead(seeded, "Aranya University")
    deal = create_deal(seeded, lead.id, DealData(title="Starter lab"), riya)
    items = list(seeded.scalars(select(PriceItem).order_by(PriceItem.id)))
    with pytest.raises(ProcurementError, match="no partner price"):
        lines_from_price_items(seeded, [(items[1].id, 1)], "partner")
    lines = lines_from_price_items(seeded, [(items[0].id, 1), (items[1].id, 1)], "institute")
    assert [line.section for line in lines] == ["software", "hardware"]
    on = date(2026, 10, 20)
    q1 = create_quote(seeded, deal.id, QuoteDraft(lines=lines, ref_code="ALU"), riya, on=on)
    seeded.commit()
    assert q1.number == "DEMO/ALU/2026-27/001" and q1.version == 1
    assert q1.attention == "Vice Chancellor" and q1.greeting == "Dear Dr. Menon,"
    assert q1.prepared_by == "Riya Sharma, Head of Academic Partnerships"
    assert q1.valid_until == on + timedelta(days=30)
    messages = mark_quote_sent(seeded, q1.id, riya, on=on)
    seeded.commit()
    assert lead.stage.name == "Sent Quotation"
    assert float(deal.value_inr) == 3550000.0  # excluding GST
    reminder = seeded.scalar(select(Activity).where(Activity.notes == f"quote:{q1.id}"))
    assert reminder is not None and reminder.status == "planned"
    assert any("Reminder" in m for m in messages)

    q2 = revise_quote(seeded, q1.id, QuoteDraft(lines=lines, discount_pct=Decimal("5")), riya, on=on)
    assert q2.number == q1.number and q2.version == 2
    mark_quote_sent(seeded, q2.id, riya, on=on)
    seeded.commit()
    assert seeded.get(Quote, q1.id).status == "superseded"
    total, weighted = pipeline_value(quotes_frame(seeded, today=on))
    assert total == pytest.approx(3550000 * 0.95)

    accept_quote(seeded, q2.id, date(2026, 11, 1), riya)
    seeded.commit()
    assert deal.status == "won" and lead.stage.outcome == "won"
    deliveries = {item.product.name: item.expected_delivery_on for item in deal.line_items}
    assert deliveries["NetSim Q"] == date(2026, 11, 22)                          # simulator +21 days
    assert deliveries["Quantum Sensing Kit"] == date(2027, 1, 10)  # kit +70 days
    # Second quote in the same FY gets the next number, whatever the institution.
    other = create_deal(seeded, _lead(seeded, "Beta College").id, DealData(title="Q"), riya)
    q3 = create_quote(seeded, other.id, QuoteDraft(lines=[LineIn("Custom item", 1, Decimal("5000"))]), riya, on=on)
    assert q3.number == "DEMO/BC/2026-27/002"


def test_quote_expiry(seeded: Session) -> None:
    riya = _user(seeded, "Riya")
    lead = _lead(seeded)
    deal = create_deal(seeded, lead.id, DealData(title="X"), riya)
    q = create_quote(seeded, deal.id, QuoteDraft(lines=[LineIn("Item", 1, Decimal("1000"))], valid_days=10),
                     riya, on=date(2026, 9, 1))
    mark_quote_sent(seeded, q.id, riya, on=date(2026, 9, 1))
    assert expire_quotes(seeded, today=date(2026, 9, 20)) == 1
    assert q.status == "expired"


def test_quote_word_file_from_template(seeded: Session) -> None:
    riya = _user(seeded, "Riya")
    _prices(seeded, riya)
    lead = _lead(seeded, "Aranya University")
    deal = create_deal(seeded, lead.id, DealData(title="Lab"), riya)
    items = list(seeded.scalars(select(PriceItem).order_by(PriceItem.id)))
    lines = lines_from_price_items(seeded, [(items[0].id, 1), (items[1].id, 2)], "institute")
    q = create_quote(seeded, deal.id, QuoteDraft(lines=lines, ref_code="ALU"), riya, on=date(2026, 9, 24))
    seeded.commit()
    data, name = quote_docx(seeded, q.id)
    assert name.endswith(".docx") and "DEMO_ALU" in name
    doc = DocxDocument(io.BytesIO(data))
    text = "\n".join(p.text for p in doc.paragraphs) + "\n".join(
        c.text for t in doc.tables for r in t.rows for c in r.cells)
    assert "{" not in text  # every field filled
    assert "Aranya University" in text and "24/09/2026" in text and "Dear Dr. Menon," in text
    assert "Quantum Software licences" in text
    assert "Benchtop hardware kits" in text
    assert "₹11,00,000" in text and "₹49,00,000" in text      # 2 × 24,50,000
    assert "Total including GST" in text and "Summary" in text  # the demo shows a summary when there are 2+ sections
    assert "5. Terms" in text and "Demo · 1 Example Tech Park" not in text and "Riya Sharma" in text and "GSTIN" not in text
    heading = next(p for p in doc.paragraphs if "Software" in p.text and "licences" in p.text)
    assert any(r.underline and r.text == "Software" for r in heading.runs)  # the [bracketed] word is underlined


# ---------------------------------------------------------------- configuration editors and the team's template


def test_checklist_rules_and_flagged_words_editors(seeded: Session) -> None:
    from app.services.library import content_flags, flagged_frame, save_flagged
    from app.services.procurement import rules_frame, save_rules

    riya, kabir = _user(seeded, "Riya"), _user(seeded, "Kabir")
    rules = rules_frame(seeded)
    assert not rules["Buying route"].str.contains("GeM ").any()  # Demo is not on GeM
    added = pd.concat([rules, pd.DataFrame([{"id": None, "Buying route": "Any", "Document needed": "Signed NDA",
                                             "Only for reseller deals": True, "Active": True}])])
    with pytest.raises(PermissionError):
        save_rules(seeded, added, kabir)
    assert save_rules(seeded, added, riya) == 1
    words = pd.concat([flagged_frame(seeded), pd.DataFrame([{"id": None, "Word or phrase": "Type-1 SPDC",
                                                             "Use instead": "entangled photon-pair source",
                                                             "Active": True}])])
    assert save_flagged(seeded, words, riya) == 1
    assert content_flags(seeded, "uses a Type-1 SPDC source") == [("Type-1 SPDC", "entangled photon-pair source")]


def test_bundled_template_is_the_demo_quotation() -> None:
    import io

    from docx import Document

    from app.services.quote_docx import TEMPLATE_PATH, check_template

    data = TEMPLATE_PATH.read_bytes()
    fields = check_template(data)
    assert "{submitted_to}" in fields and "{title_keyword}" in fields
    text = "\n".join(p.text for p in Document(io.BytesIO(data)).paragraphs)
    # The demo's own fictional template: generic sections, no real company details.
    for line in ("1. Platforms and prices", "2. What is included", "3. Payment terms",
                 "4. Warranty and support", "5. Terms", "{greeting}", "{SUMMARY}"):
        assert line in text, line
    assert "GSTIN" not in text


def test_section_title_underline_split() -> None:
    from app.services.quote_docx import split_title

    assert split_title("Quantum [Software] licences") == ("Quantum ", "Software", " licences")
    assert split_title("Services") == ("Services", "", "")


def test_removing_library_links(seeded: Session) -> None:
    from app.services.library import DocumentData, add_document, attach_documents, delete_document

    riya, kabir, neha = _user(seeded, "Riya"), _user(seeded, "Kabir"), _user(seeded, "Neha")
    mine, _ = add_document(seeded, DocumentData(title="Draft deck", doc_type="Presentation",
                                                url="https://drive.google.com/file/d/one/view"), kabir)
    other, _ = add_document(seeded, DocumentData(title="Spec sheet", doc_type="Spec sheet",
                                                 url="https://drive.google.com/file/d/two/view"), kabir)
    seeded.commit()
    with pytest.raises(LibraryError):
        delete_document(seeded, mine.id, neha)  # not hers
    assert "removed" in delete_document(seeded, mine.id, kabir)  # own draft
    lead = _lead(seeded)
    activity = Activity(type_id=seeded.scalar(select(ActivityType.id).where(ActivityType.name == "Email")),
                        lead_id=lead.id, org_id=lead.org_id, title="Sent spec", status="done", direction="outbound")
    seeded.add(activity)
    seeded.flush()
    attach_documents(seeded, activity, [other.id])
    with pytest.raises(LibraryError, match="history"):
        delete_document(seeded, other.id, riya)  # shared with a lead → archive instead
