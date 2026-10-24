"""Builds a quotation as a Word file from the team's own template (letterhead, sections, terms).

The template (assets/quote_template.docx, or one uploaded on Configuration) holds {fields} that are filled here,
plus one *model* item section — heading, intro and a table with a header row, one item row and a total row — that
is copied once per section (Software kits, Hardware kits, Services). Everything else in the template (pre-requisites,
inclusions, payment terms, support, T&Cs) is left exactly as written, so the wording is edited in Word, not code.
"""

from __future__ import annotations

import base64
import binascii
import copy
import io
import re
import zipfile
from dataclasses import dataclass
from decimal import Decimal
from pathlib import Path

from docx import Document
from docx.document import Document as DocxDocument
from docx.table import Table
from docx.text.paragraph import Paragraph

from app.services.money import inr, paise

TEMPLATE_PATH = Path(__file__).resolve().parents[2] / "assets" / "quote_template.docx"
TEMPLATE_SETTING = "quote_template_b64"   # an uploaded replacement, stored in app_settings
REQUIRED_FIELDS = ("{name}", "{section_total}")
HEADING_FIELDS = ("{title_keyword}", "{section_title}")  # either form marks the model section heading
ROMAN = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII"]
MAX_TEMPLATE_BYTES = 3 * 1024 * 1024


class TemplateError(ValueError):
    """Shown to the person as-is."""


@dataclass(frozen=True)
class Section:
    key: str
    title: str
    intro: str


def money(value) -> str:
    """₹11,00,000 for whole rupees, ₹11,00,000.50 otherwise — as in the team's quotes."""
    amount = paise(value)
    return inr(amount, decimals=amount != amount.to_integral_value())


def template_bytes(settings: dict[str, str]) -> bytes:
    uploaded = settings.get(TEMPLATE_SETTING)
    if uploaded:
        try:
            return base64.b64decode(uploaded)
        except (binascii.Error, ValueError):
            pass  # fall back to the bundled template rather than failing every quote
    return TEMPLATE_PATH.read_bytes()


def check_template(data: bytes) -> list[str]:
    """Validates an uploaded template: a real .docx, not too big, with the fields the hub needs."""
    if len(data) > MAX_TEMPLATE_BYTES:
        raise TemplateError("The template is larger than 3 MB — remove large images and try again.")
    if not zipfile.is_zipfile(io.BytesIO(data)):
        raise TemplateError("That isn't a Word (.docx) file.")
    try:
        doc = Document(io.BytesIO(data))
    except Exception as error:  # noqa: BLE001 — python-docx raises several types for broken files
        raise TemplateError("Word couldn't open that file as a document.") from error
    text = _all_text(doc)
    missing = [f for f in REQUIRED_FIELDS if f not in text]
    if not any(f in text for f in HEADING_FIELDS):
        missing.append("{title_keyword}")
    if missing:
        raise TemplateError("The template is missing these fields: " + ", ".join(missing)
                            + ". Start from the current template (Download template) and keep the model item table.")
    return sorted(set(re.findall(r"\{[A-Za-z_]+\}", text)))


def _all_text(doc: DocxDocument) -> str:
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def _fill_runs(paragraph: Paragraph, values: dict[str, str]) -> None:
    # Fields sit inside single runs in the template, so replacing run by run keeps each run's formatting.
    for run in paragraph.runs:
        if "{" in run.text:
            run.text = re.sub(r"\{([A-Za-z_]+)\}", lambda m: values.get(m.group(1), m.group(0)), run.text)


def _fill_element(element, doc: DocxDocument, values: dict[str, str]) -> None:
    for p in element.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p"):
        _fill_runs(Paragraph(p, doc), values)


def _find_paragraph(doc: DocxDocument, marker: str) -> Paragraph | None:
    for p in doc.paragraphs:
        if marker in p.text:
            return p
    return None


def build_quote_docx(*, values: dict[str, str], sections: list[tuple[Section, list[dict]]],
                     summary: list[tuple[str, str, bool]] | None, template: bytes) -> bytes:
    """values: cover fields ({title}, {submitted_to}…). sections: (Section, lines) where each line has
    name, details, qty, unit_price, amount (already formatted). summary: (label, amount, bold) rows or None."""
    doc = Document(io.BytesIO(template))
    heading = _find_paragraph(doc, "{title_keyword}") or _find_paragraph(doc, "{section_title}")
    if heading is None:
        raise TemplateError("The quotation template has no model section heading ({title_keyword}).")
    intro_el = heading._p.getnext()
    table_el = intro_el.getnext()
    spacer_el = table_el.getnext()
    if table_el.tag.split("}")[1] != "tbl":
        raise TemplateError("In the template, the model item table must come right after the section intro.")
    pristine_table = copy.deepcopy(table_el)

    for number, (section, lines) in enumerate(sections):
        block = [copy.deepcopy(el) for el in (heading._p, intro_el, table_el, spacer_el)]
        for el in block:
            heading._p.addprevious(el)
        before, keyword, after = split_title(section.title)
        _fill_element(block[0], doc, {"section_no": f"{ROMAN[number]}.", "section_title": section.title,
                                      "title_before": before, "title_keyword": keyword, "title_after": after})
        _fill_element(block[1], doc, {"section_intro": section.intro})
        table = Table(block[2], doc)
        model_row, total_row = table.rows[1]._tr, table.rows[-1]._tr
        for i, line in enumerate(lines, start=1):
            row = copy.deepcopy(model_row)
            total_row.addprevious(row)
            _fill_element(row, doc, {"n": str(i), "name": line["name"], "details": line.get("details") or "",
                                     "qty": str(line["qty"]), "unit_price": line["unit_price"],
                                     "amount": line["amount"]})
            if not line.get("details"):
                # Drop the empty second paragraph so a one-line item doesn't get a blank line.
                cell_ps = [p for p in row.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p")]
                for p in cell_ps:
                    if Paragraph(p, doc).text == "" and p.getparent() is not None and len(
                            p.getparent().findall("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p")) > 1:
                        p.getparent().remove(p)
        model_row.getparent().remove(model_row)
        _fill_element(total_row, doc, {"total_label": "Total (excluding GST)", "section_total": section_total(lines)})

    for el in (heading._p, intro_el, table_el, spacer_el):
        el.getparent().remove(el)

    marker = _find_paragraph(doc, "{SUMMARY}")
    if marker is not None:
        if summary:
            table = copy.deepcopy(pristine_table)
            marker._p.addprevious(table)
            t = Table(table, doc)
            header = t.rows[0]
            for cell, text in zip(header.cells, ["", "Summary", "", "", "Amount"]):
                for p in cell.paragraphs:
                    for j, run in enumerate(p.runs):
                        run.text = text if j == 0 else ""
            item_tr, total_tr = t.rows[1]._tr, t.rows[-1]._tr
            item_tr.getparent().remove(item_tr)
            for label, amount, bold in summary:
                row = copy.deepcopy(total_tr)
                total_tr.addprevious(row)
                _fill_element(row, doc, {"total_label": label, "section_total": amount})
                if not bold:
                    for r in row.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}b"):
                        r.getparent().remove(r)
            total_tr.getparent().remove(total_tr)
        marker._p.getparent().remove(marker._p)

    for p in doc.paragraphs:
        _fill_runs(p, values)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    _fill_runs(p, values)
    doc.core_properties.title = values.get("title", "Quotation")
    doc.core_properties.comments = f"Quotation {values.get('reference', '')}"
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def split_title(title: str) -> tuple[str, str, str]:
    """'Quantum [Software] licences' → ('Quantum ', 'Software', ' licences'): the word in
    brackets is printed underlined, as in the team's quotation."""
    match = re.match(r"^(.*?)\[(.+?)\](.*)$", title or "")
    return (match.group(1), match.group(2), match.group(3)) if match else (title or "", "", "")


def section_total(lines: list[dict]) -> str:
    return money(sum((Decimal(str(line["raw_amount"])) for line in lines), Decimal("0")))
