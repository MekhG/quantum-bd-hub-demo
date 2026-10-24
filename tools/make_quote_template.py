"""Builds assets/quote_template.docx: Demo's fictional quotation template.

The quotation builder (app/services/quote_docx.py) fills the {fields} and copies the *model section* — a heading,
an intro paragraph and a 3-row item table — once per section. Everything else is ordinary Word text that an admin
can edit in Word. The weekly report reuses this file's header and footer as its letterhead.

Run from the project folder:  python tools/make_quote_template.py
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "assets" / "quote_template.docx"
MARK = ROOT / "assets" / "demo_mark.png"

BROWN = RGBColor(0x7A, 0x4A, 0x26)
INK = RGBColor(0x2B, 0x1C, 0x10)
MUTED = RGBColor(0x6B, 0x5A, 0x4C)
BEIGE = "F7F0E6"
FONT = "Calibri"


def shade(cell, hex_fill: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    fill = OxmlElement("w:shd")
    fill.set(qn("w:val"), "clear")
    fill.set(qn("w:color"), "auto")
    fill.set(qn("w:fill"), hex_fill)
    props.append(fill)


def borders(table, colour: str = "E6D8C4") -> None:
    props = table._tbl.tblPr
    edges = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), colour)
        edges.append(el)
    props.append(edges)


def run(paragraph, text: str, *, bold: bool = False, size: float = 10.5, colour: RGBColor = INK,
        underline: bool = False, italic: bool = False):
    r = paragraph.add_run(text)
    r.font.name = FONT
    r.font.size = Pt(size)
    r.font.color.rgb = colour
    r.bold, r.underline, r.italic = bold, underline, italic
    return r


def para(doc, text: str = "", **kw):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(kw.pop("after", 6))
    if text:
        run(p, text, **kw)
    return p


def heading(doc, text: str):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    run(p, text, bold=True, size=12.5, colour=BROWN)
    return p


def simple_table(doc, rows: list[list[str]], widths: list[float]):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    borders(table)
    for i, values in enumerate(rows):
        for j, value in enumerate(values):
            cell = table.rows[i].cells[j]
            cell.width = Cm(widths[j])
            cell.paragraphs[0].paragraph_format.space_after = Pt(2)
            run(cell.paragraphs[0], value, bold=i == 0, size=9.5)
            if i == 0:
                shade(cell, BEIGE)
    return table


def letterhead(doc) -> None:
    section = doc.sections[0]
    section.top_margin, section.bottom_margin = Cm(2.6), Cm(2.0)
    section.left_margin = section.right_margin = Cm(2.0)
    header = section.header.paragraphs[0]
    header.add_run().add_picture(str(MARK), width=Cm(1.1))
    run(header, "  Demo", bold=True, size=13, colour=BROWN)
    run(header, "   ·   Quantum labs for classrooms", size=9, colour=MUTED, italic=True)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run(footer, "Demo · 1 Example Tech Park, Hyderabad, Telangana · hello@demo.example\n"
                "A fictional company created for a portfolio demonstration. This is not a real offer.",
        size=8, colour=MUTED)


def build() -> None:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = FONT
    style.font.size = Pt(10.5)
    letterhead(doc)

    title = para(doc, after=10)
    run(title, "{title}", bold=True, size=18, colour=BROWN)

    cover = doc.add_table(rows=6, cols=2)
    borders(cover)
    for i, (label, field) in enumerate((("Submitted to", "{submitted_to}"), ("Attention", "{attention}"),
                                        ("Reference", "{reference}"), ("Date", "{date}"),
                                        ("Quotation validity", "{validity}"), ("Prepared by", "{prepared_by}"))):
        left, right = cover.rows[i].cells
        left.width, right.width = Cm(4.5), Cm(12.5)
        shade(left, BEIGE)
        run(left.paragraphs[0], label, bold=True, size=9.5)
        run(right.paragraphs[0], field, size=9.5)

    para(doc, after=4)
    para(doc, "{greeting}")
    p = para(doc)
    run(p, "Thank you for your interest in a quantum teaching laboratory at ")
    run(p, "{institution}")
    run(p, ". This quotation lists the platforms we discussed, with installation and training included.")

    heading(doc, "1. Platforms and prices")
    # ---- the model section: heading → intro → item table → spacer (copied once per section) ----
    model_heading = doc.add_paragraph()
    model_heading.paragraph_format.space_before = Pt(8)
    model_heading.paragraph_format.space_after = Pt(2)
    run(model_heading, "{section_no}", bold=True, size=11, colour=BROWN)
    run(model_heading, "  ", size=11)
    run(model_heading, "{title_before}", bold=True, size=11, colour=BROWN)
    run(model_heading, "{title_keyword}", bold=True, size=11, colour=BROWN, underline=True)
    run(model_heading, "{title_after}", bold=True, size=11, colour=BROWN)
    para(doc, "{section_intro}", size=9.5, colour=MUTED, after=4)

    items = doc.add_table(rows=3, cols=5)
    items.alignment = WD_TABLE_ALIGNMENT.CENTER
    borders(items)
    widths = [1.0, 9.5, 1.4, 2.6, 2.6]
    header_cells = ["#", "Description", "Qty", "Unit price", "Amount"]
    for j, text in enumerate(header_cells):
        cell = items.rows[0].cells[j]
        cell.width = Cm(widths[j])
        shade(cell, BEIGE)
        run(cell.paragraphs[0], text, bold=True, size=9.5)
    model = items.rows[1].cells
    for j, field in enumerate(["{n}", None, "{qty}", "{unit_price}", "{amount}"]):
        model[j].width = Cm(widths[j])
        if field:
            run(model[j].paragraphs[0], field, size=9.5)
    run(model[1].paragraphs[0], "{name}", bold=True, size=9.5)
    run(model[1].add_paragraph(), "{details}", size=8.5, colour=MUTED)
    total = items.rows[2].cells
    run(total[1].paragraphs[0], "{total_label}", bold=True, size=9.5)
    run(total[4].paragraphs[0], "{section_total}", bold=True, size=9.5)
    for j in range(5):
        total[j].width = Cm(widths[j])
    para(doc, after=4)  # spacer that travels with each section
    # ---- end of the model section ----

    para(doc, "{SUMMARY}", after=4)

    heading(doc, "2. What is included")
    simple_table(doc, [
        ["Included", "What it covers"],
        ["Installation", "Set-up on site (kits) or on the institution's server (simulators), with acceptance testing."],
        ["Product training", "Hands-on sessions for the faculty who will run the laboratory."],
        ["Laboratory manuals", "Step-by-step experiment guides for students and instructors."],
        ["Remote support", "Help desk for the warranty period."],
    ], [4.5, 12.5])

    heading(doc, "3. Payment terms")
    simple_table(doc, [
        ["Configuration", "Advance with the order", "Balance"],
        ["Simulators only", "30%", "70% on installation and acceptance"],
        ["Experimentation kits", "60%", "40% on delivery and acceptance"],
    ], [5.0, 4.5, 7.5])
    para(doc, "Payments by bank transfer to the account shown on the invoice.", size=9.5, after=4)

    heading(doc, "4. Warranty and support")
    para(doc, "Twelve months' warranty from acceptance. From the second year, an annual support plan is available at "
              "5%, 7%, 9% and 11% of the equipment value for years 2 to 5.", size=10)

    heading(doc, "5. Terms")
    for line in ("This quotation is valid for {valid_days} days from the date of issue.",
                 "Delivery periods run from receipt of the purchase order and the advance payment: about three weeks for "
                 "simulators and about ten weeks for experimentation kits.",
                 "GST is charged as applicable at the time of invoice.",
                 "Specifications may be adapted to the institution's needs, subject to technical feasibility."):
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        run(p, line, size=10)

    para(doc, after=8)
    para(doc, "We would be glad to arrange a demonstration for your faculty.", after=10)
    para(doc, "Yours sincerely,", after=18)
    para(doc, "{signature_name}", bold=True, after=0)
    para(doc, "{signature_title}", colour=MUTED, after=0)
    para(doc, "Demo", colour=MUTED)

    doc.core_properties.title = "Demo quotation template"
    doc.core_properties.author = "Demo"
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print("written", OUT)


if __name__ == "__main__":
    build()
